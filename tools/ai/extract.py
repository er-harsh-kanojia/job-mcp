"""Resume text extraction (PDF via pypdf, DOCX via python-docx)."""

from __future__ import annotations

import re
from pathlib import Path


class ResumeError(ValueError):
    """Resume missing / unreadable (code + human message)."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


def _clean(text: str) -> str:
    text = text.replace("\x00", "").replace(" ", " ").replace("\r", "\n")
    lines = [re.sub(r"[ \t]+", " ", ln).strip() for ln in text.split("\n")]
    out: list[str] = []
    for ln in lines:
        if not ln and (not out or not out[-1]):
            continue  # collapse blank runs
        out.append(ln)
    return "\n".join(out).strip()


def _pdf_text(path: Path) -> tuple[str, int]:
    from pypdf import PdfReader
    try:
        reader = PdfReader(str(path))
        pages = [(p.extract_text() or "") for p in reader.pages]
    except Exception as exc:
        raise ResumeError("resume_unreadable", f"Could not read the PDF: {type(exc).__name__}") from exc
    return "\n".join(pages), len(pages)


def _docx_text(path: Path) -> tuple[str, int]:
    import docx
    try:
        d = docx.Document(str(path))
    except Exception as exc:
        raise ResumeError("resume_unreadable", f"Could not read the DOCX: {type(exc).__name__}") from exc
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(dict.fromkeys(cells)))
    text = "\n".join(parts)
    # rough page estimate (~ 550 words per page)
    return text, max(1, round(len(text.split()) / 550))


def extract_resume(path: str | Path) -> dict:
    """{"text", "pages", "filename", "words"}; raises ResumeError."""
    p = Path(path)
    if not p.is_file():
        raise ResumeError("resume_missing", "No resume uploaded yet (or the file was moved). Upload a PDF or DOCX.")
    ext = p.suffix.lower()
    if ext == ".doc":
        raise ResumeError("resume_legacy_doc",
                          "Legacy .doc files can't be read. Save the resume as PDF or DOCX and upload it again.")
    if ext == ".pdf":
        raw, pages = _pdf_text(p)
    elif ext == ".docx":
        raw, pages = _docx_text(p)
    else:
        raise ResumeError("resume_type", f"Unsupported resume type '{ext}'. Upload a PDF or DOCX.")
    text = _clean(raw)
    if len(text) < 80:
        raise ResumeError("resume_empty",
                          "No text could be extracted from the resume (is it a scanned image?). "
                          "Upload a text-based PDF or a DOCX.")
    return {"text": text, "pages": pages, "filename": p.name, "words": len(text.split())}


def number_lines(text: str) -> tuple[str, dict[str, str]]:
    """Give every non-empty line a stable id (L001, L002, ...).

    Returns (numbered text for the prompt, {id: line}).
    """
    lines: dict[str, str] = {}
    out = []
    n = 0
    for ln in text.split("\n"):
        if not ln.strip():
            out.append("")
            continue
        n += 1
        lid = f"L{n:03d}"
        lines[lid] = ln
        out.append(f"{lid}: {ln}")
    return "\n".join(out), lines
